# Files/convert.py  (yangi)
import re
from pathlib import Path
from typing import Optional, Dict, Any
from PyPDF2 import PdfReader
from docx import Document
import openpyxl

## ────────── Normalization helpers (DOCX / PDF) ────────── ##
CITATION_PATTERN = re.compile(r"\[[0-9]+\]|\([^)]*\d{4}[^)]*\)")
NUMBER_PATTERN   = re.compile(r"(?<!\d)\d+(?!\d)")
SPECIALS         = re.compile(r"[^\w\u0400-\u04FF\s\.,;:!?\-'\"]")
COMPACT_SYM      = re.compile(r"([|\.\-])\1+")

def normalize_text(text: str) -> str:
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"^\s*\d+\s*$\n?", "", t, flags=re.MULTILINE)           # page nos
    t = CITATION_PATTERN.sub("", t)
    t = re.sub(r"(?:https?://\S+|www\.\S+)(?:\s+\S+)?", "", t)          # URLs
    t = re.sub(r"-\s*\n\s*", "", t)                                    # hyphen break
    t = re.sub(r"(\w)\s*-\s*(\w)", r"\1-\2", t)                        # inline hyphen
    t = NUMBER_PATTERN.sub("", t)                                      # 1 2 3
    t = re.sub(r"[=_\-]{5,}", "", t)                                   # ______ -----
    t = re.sub(r"\b[IVXLCDM]+\s+BOB\.?\b", "", t, flags=re.I)          # I BOB
    t = re.sub(r"\b\d+(?:\.\d+)*\.?\s*§?", "", t)                      # 1.2.1
    t = re.sub(r"\.{2,}", "", t)                                       # .....  ..
    t = re.sub(r"\bNukus[-–]\d{4}\s*yil\b", "", t, flags=re.I)
    t = re.sub(r"«Қарақалпақетан» баспасы,\s*\d{4}\.", "", t)
    t = re.sub(r"\b[IVXLCDM]+\b", "", t)                               # stand‑alone Roman
    t = re.sub(r"\bтом\b", "", t, flags=re.I)
    t = SPECIALS.sub("", t)
    t = COMPACT_SYM.sub(r"\1", t)                                      # ... → .  || → |
    t = re.sub(r"\s+", " ", t)
    t = "\n".join(l.strip() for l in t.split("\n"))
    t = re.sub(r"\n{2,}", "\n\n", t)
    return t.strip()

# ────────── Statistics ────────── ##
def text_stats(text: str):
    tokens = re.findall(r"\b\w+\b", text.lower())
    return len(tokens), len(set(tokens)), sum(1 for s in re.split(r"[.!?]+", text) if s.strip())

# ────────── Main extractor / cleaner ──────────##
def extract_and_clean_text(file_path: str, ext: Optional[str] = None) -> Dict[str, Any]:
    """
    Return dict with keys:
        text            → cleaned text (str)
        token_count     → int | None
        vocab_count     → int | None
        sentence_count  → int | None
    """
    path = Path(file_path)
    ext = (ext or path.suffix.lstrip('.')).lower()

    # ---------- Extract ----------##
    if ext == "pdf":
        reader = PdfReader(path)
        raw = "\n".join(page.extract_text() or "" for page in reader.pages)

    elif ext == "docx":
        raw = "\n".join(p.text for p in Document(path).paragraphs)

    elif ext in {"xls", "xlsx"}:
        wb = openpyxl.load_workbook(path, data_only=True)
        rows = [
            "|".join(str(cell) for cell in row if cell is not None)
            for sh in wb.worksheets
            for row in sh.iter_rows(values_only=True)
        ]
        raw = "\n".join(rows)

    elif ext == "txt":
        raw = path.read_text(encoding="utf-8")

    else:
        raise ValueError(f"Qo‘llab‑quvvatlanmaydigan format: .{ext}")

    # ---------- Clean ----------
    if ext in {"pdf", "docx", "txt"}:
        cleaned = normalize_text(raw)
        tok, voc, sent = text_stats(cleaned)
    else:  # excel
        cleaned = raw
        tok = voc = sent = None

    return {
        "text": cleaned,
        "token_count": tok,
        "vocab_count": voc,
        "sentence_count": sent,
    }